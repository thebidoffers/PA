import json
from collections.abc import Mapping
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument

from services.placeholder_service import extract_placeholders_from_docx

SCHEMA_PATH = Path("prompts/input_schema_talabat.json")
SUPPORTED_SCHEMA_ID = "talabat_v1"

FIELD_DEFAULTS: dict[str, Any] = {
    "issuer.name": "",
    "offer.offer_shares": None,
    "offer.percentage_offered": None,
    "offer.nominal_value_per_share_aed": None,
    "offer.price_range_low_aed": None,
    "offer.price_range_high_aed": None,
    "key_dates": "",
    "business_description": "",
    "risk_factors": [],
    "tranche_1.min_subscription_aed": None,
    "tranche_1.increment_aed": None,
    "tranche_2.min_subscription_aed": None,
}

DERIVED_PLACEHOLDER_DEPENDENCIES: dict[str, list[str]] = {
    "offer.price_range": ["offer.price_range_low_aed", "offer.price_range_high_aed"],
    "offer.nominal_value_per_share": ["offer.nominal_value_per_share_aed"],
    "offer.size": ["offer.offer_shares"],
}


def _deep_get(data: Mapping[str, Any], path: str) -> Any:
    current: Any = data
    for part in path.split("."):
        if not isinstance(current, Mapping) or part not in current:
            return None
        current = current[part]
    return current


def _deep_set(data: dict[str, Any], path: str, value: Any) -> None:
    current = data
    parts = path.split(".")
    for part in parts[:-1]:
        if part not in current or not isinstance(current[part], dict):
            current[part] = {}
        current = current[part]
    current[parts[-1]] = value


def load_schema() -> dict[str, Any]:
    if not SCHEMA_PATH.exists():
        raise FileNotFoundError(f"Talabat schema file is missing: {SCHEMA_PATH}")
    schema = json.loads(SCHEMA_PATH.read_text())
    if schema.get("schema_id") != SUPPORTED_SCHEMA_ID:
        raise ValueError(f"Unsupported schema file: {schema.get('schema_id')}")
    return schema


def extract_template_placeholders(template_path: str | Path) -> list[str]:
    document = DocxDocument(str(template_path))
    return extract_placeholders_from_docx(document)


def build_template_form_spec(template_placeholders: list[str], schema: Mapping[str, Any]) -> dict[str, Any]:
    schema_fields = {field["path"]: field for field in schema["fields"]}

    requested_paths: set[str] = set()
    for placeholder in template_placeholders:
        if placeholder in schema_fields:
            requested_paths.add(placeholder)
            continue
        requested_paths.update(DERIVED_PLACEHOLDER_DEPENDENCIES.get(placeholder, []))

    requested_fields = [schema_fields[path] for path in schema_fields if path in requested_paths]
    required_paths = {field["path"] for field in requested_fields if field.get("required")}

    return {
        "schema_id": schema["schema_id"],
        "fields": requested_fields,
        "required_paths": sorted(required_paths),
        "requested_paths": sorted(requested_paths),
    }


def build_raw_inputs_payload(
    schema_id: str,
    project_id: int,
    template_id: int,
    source_document_id: int | None,
    use_template_as_source: bool,
    field_values: Mapping[str, Any],
) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "schema_id": schema_id,
        "project_id": project_id,
        "template_id": template_id,
        "source_document_id": source_document_id,
        "use_template_as_source": use_template_as_source,
    }

    for path, default in FIELD_DEFAULTS.items():
        value = field_values.get(path, default)
        if path == "risk_factors" and isinstance(value, str):
            value = [line.strip() for line in value.splitlines() if line.strip()]
        _deep_set(payload, path, value)

    _deep_set(payload, "offer.currency", "AED")
    return payload


def find_unresolved_template_placeholders(
    template_placeholders: list[str],
    rendered_map: Mapping[str, str],
) -> list[str]:
    unresolved: list[str] = []
    for placeholder in template_placeholders:
        value = rendered_map.get(placeholder)
        if value is None or str(value).startswith("[[MISSING:"):
            unresolved.append(placeholder)
    return sorted(set(unresolved))


def validate_required_paths(
    required_paths: list[str],
    raw_payload: Mapping[str, Any],
    rendered_map: Mapping[str, str],
) -> list[str]:
    errors: list[str] = []
    for path in required_paths:
        rendered_value = rendered_map.get(path)
        if rendered_value is not None and not str(rendered_value).startswith("[[MISSING:"):
            continue

        raw_value = _deep_get(raw_payload, path)
        if raw_value is None:
            errors.append(f"{path} is required.")
            continue
        if isinstance(raw_value, str) and not raw_value.strip():
            errors.append(f"{path} is required.")
            continue
    return errors
