import json
from pathlib import Path
from typing import Any

import streamlit as st
from docx import Document as DocxDocument

from db.init_db import init_db
from db.session import SessionLocal
from models import Document, ProspectusProject, Template
from services.document_service import extract_preview_and_outline
from services.generation_service import generate_draft_docx
from services.normalization_service import normalize_inputs
from services.placeholder_service import extract_placeholders_from_docx

SCHEMA_PATH = Path("prompts/input_schema_talabat.json")

init_db()

st.title("AUTO GENERATION")
st.caption("Validated inputs + assembly + generation runs.")

session = SessionLocal()
try:
    templates = session.query(Template).order_by(Template.created_at.desc()).all()
    projects = session.query(ProspectusProject).order_by(ProspectusProject.created_at.desc()).all()
finally:
    session.close()

if not templates or not projects:
    st.info("You need at least one template and one project to start generation.")
    st.stop()

if not SCHEMA_PATH.exists():
    st.error(f"Talabat schema file is missing: {SCHEMA_PATH}")
    st.stop()

schema = json.loads(SCHEMA_PATH.read_text())
field_meta = {field["path"]: field for field in schema["fields"]}


def _field_help(path: str) -> str:
    field = field_meta[path]
    return f"{field['help_text']} Example: {field['example']}"


def _build_inputs_payload(template_id: int, project_id: int, source_document_id: int | None, use_template_as_source: bool) -> dict[str, Any]:
    risk_lines = [line.strip() for line in st.session_state.get("risk_factors_input", "").splitlines() if line.strip()]
    return {
        "schema_id": schema["schema_id"],
        "issuer": {"name": st.session_state.get("issuer_name", "").strip()},
        "offer": {
            "offer_shares": st.session_state.get("offer_offer_shares"),
            "percentage_offered": st.session_state.get("offer_percentage_offered"),
            "nominal_value_per_share_aed": st.session_state.get("offer_nominal_value_per_share_aed"),
            "price_range_low_aed": st.session_state.get("offer_price_range_low_aed"),
            "price_range_high_aed": st.session_state.get("offer_price_range_high_aed"),
            "currency": "AED",
        },
        "key_dates": st.session_state.get("key_dates", "").strip(),
        "business_description": st.session_state.get("business_description", "").strip(),
        "risk_factors": risk_lines,
        "tranche_1": {
            "min_subscription_aed": st.session_state.get("tranche_1_min_subscription_aed"),
            "increment_aed": st.session_state.get("tranche_1_increment_aed"),
        },
        "tranche_2": {
            "min_subscription_aed": st.session_state.get("tranche_2_min_subscription_aed"),
        },
        "source_document_id": source_document_id,
        "use_template_as_source": use_template_as_source,
        "template_id": template_id,
        "project_id": project_id,
    }


template = st.selectbox("Template", options=templates, format_func=lambda t: f"#{t.id} {t.name} ({t.status})")
project = st.selectbox("Project", options=projects, format_func=lambda p: f"#{p.id} {p.name}")

template_doc = DocxDocument(template.file_path)
template_placeholders = extract_placeholders_from_docx(template_doc)
if not template_placeholders:
    st.error("Selected template has placeholder_count=0. Generation is blocked. Use Templates → Auto-Parameterize from Source Prospectus first.")
    st.stop()

session = SessionLocal()
try:
    project_documents = (
        session.query(Document)
        .filter(Document.project_id == project.id)
        .order_by(Document.created_at.desc())
        .all()
    )
finally:
    session.close()

use_template_as_source = False
source_document = None
source_for_preview_path = template.file_path

if project_documents:
    source_document = st.selectbox(
        "Source document",
        options=project_documents,
        format_func=lambda d: f"#{d.id} {d.doc_type} v{d.version} ({d.file_name})",
    )
    source_for_preview_path = source_document.file_path
else:
    st.warning("Selected project has no documents.")
    use_template_as_source = st.checkbox(
        "Use selected template as source document for preview (skip project upload)",
        value=False,
    )
    if not use_template_as_source:
        st.info("Enable the checkbox above to preview from the template and generate without a project upload.")

if source_document is not None or use_template_as_source:
    preview_data = extract_preview_and_outline(source_for_preview_path)

    st.divider()
    left, right = st.columns(2)
    with left:
        st.subheader("Extracted Text Preview")
        st.text_area("Preview", value=preview_data["preview"] or "(empty text)", height=260, disabled=True)

    with right:
        st.subheader("Outline JSON Preview")
        st.code(json.dumps(preview_data["outline"], indent=2), language="json")

st.divider()
st.subheader("Generation Inputs")

st.text_input("Issuer Name", key="issuer_name", help=_field_help("issuer.name"))
st.number_input(
    "Offer Shares",
    min_value=0,
    step=1,
    key="offer_offer_shares",
    help=_field_help("offer.offer_shares"),
)
st.number_input(
    "Offer Price Range per Offer Share (AED) - Low",
    min_value=0.0,
    step=0.01,
    format="%.2f",
    key="offer_price_range_low_aed",
    help=_field_help("offer.price_range_low_aed"),
)
st.number_input(
    "Offer Price Range per Offer Share (AED) - High",
    min_value=0.0,
    step=0.01,
    format="%.2f",
    key="offer_price_range_high_aed",
    help=_field_help("offer.price_range_high_aed"),
)
st.number_input(
    "Nominal Value per Share (AED)",
    min_value=0.0,
    step=0.01,
    format="%.2f",
    key="offer_nominal_value_per_share_aed",
    help=_field_help("offer.nominal_value_per_share_aed"),
)
st.number_input(
    "Percentage Offered",
    min_value=0.0,
    max_value=100.0,
    step=0.01,
    format="%.2f",
    key="offer_percentage_offered",
    help=_field_help("offer.percentage_offered"),
)
st.text_input("Key Dates", key="key_dates", help=_field_help("key_dates"))
st.text_area("Business Description", key="business_description", help=_field_help("business_description"))
st.text_area("Risk Factors (one per line)", key="risk_factors_input", help=_field_help("risk_factors"))
st.number_input(
    "Tranche 1 Minimum Subscription (AED)",
    min_value=0,
    step=1,
    key="tranche_1_min_subscription_aed",
    help=_field_help("tranche_1.min_subscription_aed"),
)
st.number_input(
    "Tranche 1 Increment (AED)",
    min_value=0,
    step=1,
    key="tranche_1_increment_aed",
    help=_field_help("tranche_1.increment_aed"),
)
st.number_input(
    "Tranche 2 Minimum Subscription (AED)",
    min_value=0,
    step=1,
    key="tranche_2_min_subscription_aed",
    help=_field_help("tranche_2.min_subscription_aed"),
)

source_document_id = source_document.id if source_document is not None else None
inputs_payload = _build_inputs_payload(template.id, project.id, source_document_id, use_template_as_source)
_, rendered_preview, _ = normalize_inputs(schema["schema_id"], inputs_payload)

st.subheader("Normalized Preview")
st.json(
    {
        "Offer Shares": rendered_preview.get("offer.offer_shares"),
        "Offer Price Range per Offer Share": rendered_preview.get("offer.price_range"),
        "Nominal Value": rendered_preview.get("offer.nominal_value_per_share"),
        "Percentage Offered": rendered_preview.get("offer.percentage_offered"),
    }
)

confirm_disclaimer = st.checkbox(
    "I confirm all facts are verified and that missing facts are marked as TBD or [[MISSING: field]]."
)

if st.button("Generate"):
    required_errors: list[str] = []

    issuer_name = st.session_state.get("issuer_name", "").strip()
    offer_shares_value = st.session_state.get("offer_offer_shares", 0)
    low = st.session_state.get("offer_price_range_low_aed", 0.0)
    high = st.session_state.get("offer_price_range_high_aed", 0.0)

    if not issuer_name:
        required_errors.append("issuer.name is required.")
    if offer_shares_value is None or int(offer_shares_value) <= 0:
        required_errors.append("Offer Shares must be greater than 0.")
    if low is None or high is None or float(low) >= float(high):
        required_errors.append("Offer Price Range requires low < high.")
    if source_document is None and not use_template_as_source:
        required_errors.append("Select a source document or enable 'Use selected template as source document'.")

    if required_errors:
        for error in required_errors:
            st.error(error)
    elif not confirm_disclaimer:
        st.error("You must confirm the disclaimer before generating.")
    else:
        result = generate_draft_docx(project.id, template.id, inputs_payload)

        output_path = Path(result["output_path"])
        st.success(f"Generation run #{result['generation_run_id']} completed.")
        if result["missing_fields"]:
            st.warning("Missing fields (template-driven): " + ", ".join(result["missing_fields"]))

        with output_path.open("rb") as generated_file:
            st.download_button(
                label="Download Generated Draft DOCX",
                data=generated_file.read(),
                file_name=output_path.name,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key=f"generated_download_{result['document_id']}",
            )
