import json
from pathlib import Path

import streamlit as st
from docx import Document as DocxDocument

from db.init_db import init_db
from db.session import SessionLocal
from models import Document, ProspectusProject, Template
from services.file_service import save_uploaded_file
from services.parameterization_service import parameterize_template_from_source
from services.placeholder_service import extract_placeholders_from_docx
from services.prospectus_analysis_service import analyze_prospectus, save_analysis

init_db()

st.title("TEMPLATES")
st.caption("Manage your template library (upload, status, preview, inspection, and parameterization).")

with st.form("template_upload_form"):
    st.subheader("Upload Template")
    template_name = st.text_input("Template name", placeholder="e.g., Standard Prospectus Template v1")
    status = st.selectbox("Status", options=["draft", "approved"], index=0)
    uploaded_template = st.file_uploader("Upload DOCX template", type=["docx"])
    submitted = st.form_submit_button("Save Template")

if submitted:
    if not template_name.strip() or uploaded_template is None:
        st.error("Template name and DOCX file are required.")
    else:
        safe_name = f"{template_name.strip().replace(' ', '_')}_{uploaded_template.name}"
        file_path, file_sha256 = save_uploaded_file(uploaded_template, "storage/templates", safe_name)

        session = SessionLocal()
        try:
            template = Template(
                name=template_name.strip(),
                status=status,
                sha256=file_sha256,
                file_path=str(file_path),
            )
            session.add(template)
            session.commit()
            st.success("Template saved.")
        finally:
            session.close()

st.divider()
st.subheader("Template Library")

session = SessionLocal()
try:
    templates = session.query(Template).order_by(Template.created_at.desc()).all()
    projects = session.query(ProspectusProject).order_by(ProspectusProject.created_at.desc()).all()
finally:
    session.close()

if not templates:
    st.info("No templates available yet.")
else:
    table_data = [
        {
            "ID": t.id,
            "Name": t.name,
            "Status": t.status,
            "Version": t.version,
            "SHA256": t.sha256,
            "Path": t.file_path,
            "Created At": t.created_at,
        }
        for t in templates
    ]
    st.dataframe(table_data, use_container_width=True)

    st.markdown("#### Downloads")
    for t in templates:
        template_path = Path(t.file_path)
        if template_path.exists():
            with template_path.open("rb") as f:
                st.download_button(
                    label=f"Download Template #{t.id}: {t.name}",
                    data=f.read(),
                    file_name=template_path.name,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    key=f"template_download_{t.id}",
                )
        else:
            st.warning(f"File not found for template #{t.id}: {t.file_path}")

    st.divider()
    st.subheader("Inspect Template")

    inspect_template = st.selectbox(
        "Choose template to inspect",
        options=templates,
        format_func=lambda t: f"#{t.id} {t.name} (v{t.version})",
        key="inspect_template_select",
    )
    if st.button("Inspect Template", key="inspect_template_button"):
        inspect_path = Path(inspect_template.file_path)
        if not inspect_path.exists():
            st.error(f"Template file not found: {inspect_template.file_path}")
        else:
            doc = DocxDocument(str(inspect_path))
            placeholders = extract_placeholders_from_docx(doc)
            st.write({"placeholder_count": len(placeholders), "placeholders": placeholders})
            if not placeholders:
                st.warning("No placeholders found. This is likely a static prospectus.")

st.divider()
st.subheader("Auto-Parameterize from Source Prospectus")

if not projects:
    st.info("Create a project and upload source documents before running auto-parameterization.")
else:
    selected_project = st.selectbox(
        "Project",
        options=projects,
        format_func=lambda p: f"#{p.id} {p.name}",
        key="param_project",
    )

    session = SessionLocal()
    try:
        project_docs = (
            session.query(Document)
            .filter(Document.project_id == selected_project.id, Document.file_name.ilike("%.docx"))
            .order_by(Document.created_at.desc())
            .all()
        )
    finally:
        session.close()

    if not project_docs:
        st.warning("No DOCX source documents found for this project.")
    else:
        selected_source = st.selectbox(
            "Source DOCX document",
            options=project_docs,
            format_func=lambda d: f"#{d.id} {d.doc_type} v{d.version} ({d.file_name})",
            key="param_source_doc",
        )

        allow_source_as_base = st.checkbox("Use source document itself as base template", value=False)
        base_template = None
        if not allow_source_as_base:
            if templates:
                base_template = st.selectbox(
                    "Base template",
                    options=templates,
                    format_func=lambda t: f"#{t.id} {t.name} (v{t.version})",
                    key="param_base_template",
                )
            else:
                st.warning("No templates available. Enable source-as-base or upload a template.")

        st.markdown("#### Deterministic Matching Inputs")
        issuer_name = st.text_input("issuer.name (required)", key="param_issuer_name")
        offer_shares = st.number_input("offer.offer_shares (required)", min_value=0, step=1, key="param_offer_shares")
        percentage_offered = st.number_input(
            "offer.percentage_offered (optional)", min_value=0.0, max_value=100.0, step=0.01, format="%.2f", key="param_percentage_offered"
        )
        nominal_value = st.number_input(
            "offer.nominal_value_per_share_aed (optional)", min_value=0.0, step=0.01, format="%.2f", key="param_nominal_value"
        )
        price_low = st.number_input(
            "offer.price_range_low_aed (optional)", min_value=0.0, step=0.01, format="%.2f", key="param_price_low"
        )
        price_high = st.number_input(
            "offer.price_range_high_aed (optional)", min_value=0.0, step=0.01, format="%.2f", key="param_price_high"
        )

        if st.button("Run Auto-Parameterize", key="run_auto_parameterize"):
            errors: list[str] = []
            if not issuer_name.strip():
                errors.append("issuer.name is required.")
            if int(offer_shares) <= 0:
                errors.append("offer.offer_shares must be greater than 0.")
            if not allow_source_as_base and base_template is None:
                errors.append("Select a base template or enable source document as base.")
            if (price_low > 0 or price_high > 0) and float(price_low) >= float(price_high):
                errors.append("offer.price_range requires low < high when provided.")

            if errors:
                for error in errors:
                    st.error(error)
            else:
                source_path = selected_source.file_path

                if allow_source_as_base:
                    file_bytes = Path(source_path).read_bytes()
                    session = SessionLocal()
                    try:
                        source_as_template = Template(
                            name=f"Source Template {selected_source.file_name}",
                            status="draft",
                            sha256=selected_source.sha256,
                            file_path=source_path,
                            metadata_json=json.dumps({"derived_from_document_id": selected_source.id}),
                        )
                        session.add(source_as_template)
                        session.commit()
                        session.refresh(source_as_template)
                        base_template_id = source_as_template.id
                    finally:
                        session.close()
                else:
                    base_template_id = base_template.id

                inputs = {
                    "issuer": {"name": issuer_name.strip()},
                    "offer": {
                        "offer_shares": int(offer_shares),
                        "percentage_offered": float(percentage_offered) if float(percentage_offered) > 0 else None,
                        "nominal_value_per_share_aed": float(nominal_value) if float(nominal_value) > 0 else None,
                        "price_range_low_aed": float(price_low) if float(price_low) > 0 else None,
                        "price_range_high_aed": float(price_high) if float(price_high) > 0 else None,
                    },
                }

                analysis = analyze_prospectus(source_path, issuer_name=issuer_name.strip(), offer_shares=int(offer_shares))
                analysis_id = save_analysis(selected_project.id, selected_source.id, analysis)

                result = parameterize_template_from_source(
                    source_docx_path=source_path,
                    inputs=inputs,
                    base_template_id=base_template_id,
                    source_document_id=selected_source.id,
                    project_id=selected_project.id,
                )
                st.success(
                    f"Auto-parameterization complete. New template #{result['template_id']} created. Analysis #{analysis_id} saved."
                )
                st.json(result["parameterization_report"])
                st.caption(f"Placeholder count: {result['parameterization_report']['placeholder_count']}")
